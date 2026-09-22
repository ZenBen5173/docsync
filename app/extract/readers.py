"""File -> Document readers. Each keeps a precise source location per row:

  txt   {"kind": "line", "line": 7}
  pdf   {"kind": "bbox", "page": 0, "bbox": [x0, y0, x1, y1]}   (PDF points)
  xlsx  {"kind": "cell", "sheet": "BL", "cell": "B7", "row": 7}
  docx  {"kind": "para", "para": 2} | {"kind": "table", "table": 0, "row": 3}

Readers never raise: a broken file comes back as readable=False with the error.
"""
import hashlib
import io
import re
import threading
import zipfile
from pathlib import PurePosixPath

from app.extract.model import Document, Row

_ocr_engine = None
_ocr_lock = threading.Lock()
_ocr_cache: dict[str, list] = {}        # source-file hash + page -> OCR result, so previews of a scan do not redo the work


def _ocr_files(key: str):
    """Where an OCR result lives on disk: the writable cache first, then the read-only copy shipped with a
    hosted build (deploy/ocr_cache), which has no OCR engine of its own."""
    from app import config
    return [config.VAR_DIR / "ocr_cache" / f"{key}.json", config.ROOT / "deploy" / "ocr_cache" / f"{key}.json"]


def _ocr_load(key: str):
    import json
    for f in _ocr_files(key):
        try:
            if f.exists():
                return json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            continue
    return None


def _ocr_store(key: str, result: list):
    import json
    try:
        f = _ocr_files(key)[0]
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(json.dumps(result), encoding="utf-8")
    except Exception:
        pass                                                # a read-only disk must never break a check


def read_document(path: str, data: bytes) -> Document:
    ext = PurePosixPath(path.replace("\\", "/")).suffix.lower().lstrip(".")
    fmt = _sniff(ext, data)
    doc = Document(path=path, fmt=fmt)
    if not data or not data.strip():
        doc.readable, doc.error = False, "file is empty"
        return doc
    try:
        {"txt": _read_txt, "pdf": _read_pdf, "xlsx": _read_xlsx,
         "docx": _read_docx, "image": _read_image}.get(fmt, _read_txt)(doc, data)
    except Exception as e:  # corrupted / truncated / password protected ...
        doc.readable, doc.error = False, f"{type(e).__name__}: {str(e)[:200]}"
        return doc
    if doc.readable and not any(r.text.strip() for r in doc.rows):
        doc.readable, doc.error = False, "no text could be extracted"
    return doc


def _sniff(ext: str, data: bytes) -> str:
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
            if any(n.startswith("xl/") for n in names):
                return "xlsx"
            if any(n.startswith("word/") for n in names):
                return "docx"
        except zipfile.BadZipFile:
            pass
        return ext if ext in ("xlsx", "docx") else "unknown"
    if head.startswith((b"\x89PNG", b"\xff\xd8")):
        return "image"
    return {"text": "txt", "csv": "txt", "jpg": "image", "jpeg": "image", "png": "image"}.get(ext, ext or "txt")


# ---------------------------------------------------------------- txt
def _read_txt(doc: Document, data: bytes):
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="replace")
    # binary garbage check: a text file should be overwhelmingly printable
    sample = text[:4000]
    bad = sum(1 for ch in sample if ch == "�" or (ord(ch) < 32 and ch not in "\r\n\t"))
    if sample and bad / len(sample) > 0.10:
        doc.readable, doc.error = False, "file content is not readable text"
        return
    for i, line in enumerate(text.replace("\r\n", "\n").split("\n"), start=1):
        doc.rows.append(Row(text=line.rstrip(), loc={"kind": "line", "line": i},
                            indent=bool(re.match(r"^(\s{2,}|\t)\S", line))))


# ---------------------------------------------------------------- pdf
def _read_pdf(doc: Document, data: bytes):
    import pymupdf
    pdf = pymupdf.open(stream=data, filetype="pdf")
    if pdf.needs_pass:
        doc.readable, doc.error = False, "PDF is password protected"
        return
    doc.meta["pages"] = len(pdf)
    total_words = 0
    for pno, page in enumerate(pdf):
        words = page.get_text("words")
        total_words += len(words)
        doc.rows.extend(_rows_from_words(words, pno))
    if total_words == 0:
        # image-only scan: fall back to OCR so a reviewer still gets proposed values
        doc.method = "ocr"
        doc.meta["scanned"] = True
        confs = []
        source = hashlib.sha1(data).hexdigest()
        for pno, page in enumerate(pdf):
            key = f"{source}-p{pno}"
            # the page is only rendered when the text is not already known: the hosted copy never needs to
            png = None if _ocr_known(key) else page.get_pixmap(dpi=200).tobytes("png")
            rows, conf = _ocr_rows(png, pno, scale=72 / 200, key=key)
            doc.rows.extend(rows)
            confs.extend(conf)
        doc.ocr_confidence = round(sum(confs) / len(confs), 3) if confs else 0.0


def _rows_from_words(words, pno: int, seg_gap: float = 10.0) -> list[Row]:
    if not words:
        return []
    words = sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    lines: list[list] = []
    for w in words:
        yc = (w[1] + w[3]) / 2
        if lines and abs(yc - lines[-1][0]) <= max(3.0, (w[3] - w[1]) * 0.45):
            lines[-1][1].append(w)
        else:
            lines.append([yc, [w]])
    margin = min(w[0] for w in words)
    rows = []
    for _, ws in lines:
        ws.sort(key=lambda w: w[0])
        segs = [[ws[0]]]
        for prev, cur in zip(ws, ws[1:]):
            (segs.append([cur]) if cur[0] - prev[2] > seg_gap else segs[-1].append(cur))
        seg_txt = [" ".join(w[4] for w in s) for s in segs]
        bbox = _union(ws)
        row = Row(text="   ".join(seg_txt), loc={"kind": "bbox", "page": pno, "bbox": bbox},
                  indent=ws[0][0] > margin + 40,
                  words=[(w[4], [round(w[0], 1), round(w[1], 1), round(w[2], 1), round(w[3], 1)]) for w in ws])
        if len(segs) >= 2:
            row.label = seg_txt[0]
            row.value = " ".join(seg_txt[1:])
            row.value_loc = {"kind": "bbox", "page": pno, "bbox": _union([w for s in segs[1:] for w in s])}
        rows.append(row)
    return rows


def _union(ws) -> list[float]:
    return [round(min(w[0] for w in ws), 1), round(min(w[1] for w in ws), 1),
            round(max(w[2] for w in ws), 1), round(max(w[3] for w in ws), 1)]


# ---------------------------------------------------------------- OCR
def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_known(key: str) -> bool:
    return key in _ocr_cache or any(f.exists() for f in _ocr_files(key))


def _ocr_rows(png: bytes | None, pno: int, scale: float = 1.0, key: str | None = None):
    # One OCR job at a time, and one engine per process. Without the lock the pipeline's workers built
    # several engines at once and ran inference in parallel: ~0.5 GiB peak EACH on a 200-dpi page, enough
    # to get a small container killed. Serialised, the peak is bounded whatever the number of visitors.
    key = key or hashlib.sha1(png).hexdigest()
    try:
        with _ocr_lock:
            result = _ocr_cache.get(key)
            if result is None:
                result = _ocr_load(key)
                if result is None:
                    result, _ = _get_ocr()(png)
                    # plain lists and floats, so the result can be written to disk and read back unchanged
                    result = [[[[float(x), float(y)] for x, y in box], str(text), float(conf)] for box, text, conf in (result or [])]
                    _ocr_store(key, result)
                if len(_ocr_cache) >= 64:                   # small and bounded: a page's text boxes, not images
                    _ocr_cache.pop(next(iter(_ocr_cache)))
                _ocr_cache[key] = result or []
    except Exception:
        return [], []
    rows, confs = [], []
    for box, text, conf in (result or []):
        xs, ys = [p[0] for p in box], [p[1] for p in box]
        bbox = [round(min(xs) * scale, 1), round(min(ys) * scale, 1),
                round(max(xs) * scale, 1), round(max(ys) * scale, 1)]
        rows.append(Row(text=text, loc={"kind": "bbox", "page": pno, "bbox": bbox}))
        confs.append(float(conf))
    rows.sort(key=lambda r: (r.loc["bbox"][1], r.loc["bbox"][0]))
    return rows, confs


def _read_image(doc: Document, data: bytes):
    doc.method = "ocr"
    doc.meta["scanned"] = True
    rows, confs = _ocr_rows(data, 0)
    doc.rows = rows
    doc.ocr_confidence = round(sum(confs) / len(confs), 3) if confs else 0.0


# ---------------------------------------------------------------- xlsx
def _read_xlsx(doc: Document, data: bytes):
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    doc.meta["sheets"] = wb.sheetnames
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            cells = [c for c in r if c.value is not None and str(c.value).strip() != ""]
            if not cells:
                continue
            first = cells[0]
            loc = {"kind": "cell", "sheet": ws.title, "cell": first.coordinate, "row": first.row}
            row = Row(text="   ".join(_cell_str(c.value) for c in cells), loc=loc)
            if len(cells) >= 2:
                row.label = _cell_str(first.value)
                row.value = _cell_str(cells[1].value)
                row.value_loc = {"kind": "cell", "sheet": ws.title, "cell": cells[1].coordinate, "row": first.row}
            elif isinstance(first.value, str):
                # a label whose value cell is empty still matters (blank field)
                nxt = f"{get_column_letter(first.column + 1)}{first.row}"
                row.value_loc = {"kind": "cell", "sheet": ws.title, "cell": nxt, "row": first.row}
            doc.rows.append(row)


def _cell_str(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------- docx
def _read_docx(doc: Document, data: bytes):
    import docx as pydocx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = pydocx.Document(io.BytesIO(data))
    p_i = t_i = 0
    for child in d.element.body.iterchildren():      # keep document order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, d).text
            if text.strip():
                doc.rows.append(Row(text=text.strip(), loc={"kind": "para", "para": p_i}))
            p_i += 1
        elif tag == "tbl":
            table = Table(child, d)
            for r_i, r in enumerate(table.rows):
                cells, seen = [], set()
                for c in r.cells:                     # merged cells repeat; de-dupe
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip())
                if not any(cells):
                    continue
                loc = {"kind": "table", "table": t_i, "row": r_i}
                row = Row(text="   ".join(c.replace("\n", " / ") for c in cells), loc=loc)
                if len(cells) >= 2:
                    row.label, row.value = cells[0], cells[1]
                    row.value_loc = dict(loc, col=1)
                doc.rows.append(row)
            t_i += 1
